"""
analisis-data.py
Membaca data/log-pengujian.csv (hasil run-benchmark.js), lalu menghasilkan:
  1. data/processed/ringkasan-statistik.xlsx  -> tabel statistik deskriptif per tahap & metode
  2. data/results/perbandingan-durasi.png     -> bar chart rata-rata durasi tiap tahap (JWT vs ZKP)
  3. data/results/boxplot-total-setup.png     -> box plot sebaran waktu total call setup
  4. data/results/perbandingan-payload.png    -> bar chart ukuran payload (byte) JWT vs ZKP
"""

import os
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(BASE_DIR, 'data', 'log-pengujian.csv')

PROCESSED_DIR = os.path.join(BASE_DIR, 'data', 'processed')
RESULTS_DIR = os.path.join(BASE_DIR, 'data', 'results')
os.makedirs(PROCESSED_DIR, exist_ok=True)
os.makedirs(RESULTS_DIR, exist_ok=True)

XLSX_PATH = os.path.join(PROCESSED_DIR, 'ringkasan-statistik.xlsx')
CHART_DURASI_PATH = os.path.join(RESULTS_DIR, 'perbandingan-durasi.png')
CHART_BOXPLOT_PATH = os.path.join(RESULTS_DIR, 'boxplot-total-setup.png')
CHART_PAYLOAD_PATH = os.path.join(RESULTS_DIR, 'perbandingan-payload.png')

WARNA = {'JWT': '#2196f3', 'ZKP': '#9c27b0'}


def muat_data():
    if not os.path.exists(CSV_PATH):
        raise FileNotFoundError(
            f"Tidak ditemukan: {CSV_PATH}\n"
            f"Jalankan dulu 'node run-benchmark.js' agar data/log-pengujian.csv terisi."
        )
    df = pd.read_csv(CSV_PATH)
    df['metode'] = df['metode'].astype(str).str.upper().str.replace('-GROTH16', '', regex=False)
    df['metode'] = df['metode'].replace({'-': None})
    df['durasi_ms'] = pd.to_numeric(df['durasi_ms'], errors='coerce')
    df['ukuran_payload_byte'] = pd.to_numeric(df['ukuran_payload_byte'], errors='coerce')
    return df


def hitung_ringkasan_performa(df):
    data_sukses = df[(df['status'] == 'sukses') & df['durasi_ms'].notna()]
    ringkasan = data_sukses.groupby(['metode', 'tahap'])['durasi_ms'].agg(
        n='count', rata_rata='mean', std_dev='std', median='median',
        minimum='min', maksimum='max', p95=lambda x: x.quantile(0.95)
    ).reset_index()
    return ringkasan.round(3)


def hitung_ringkasan_payload(df):
    data_sukses = df[(df['status'] == 'sukses') & df['ukuran_payload_byte'].notna()]
    ringkasan = data_sukses.groupby(['metode', 'tahap'])['ukuran_payload_byte'].agg(
        n='count', rata_rata='mean', minimum='min', maksimum='max'
    ).reset_index().round(1)
    return ringkasan


def hitung_asr(df):
    ringkasan = df.groupby(['metode', 'tahap', 'status']).size().unstack(fill_value=0)
    for kolom in ['sukses', 'gagal']:
        if kolom not in ringkasan.columns:
            ringkasan[kolom] = 0
    ringkasan['total'] = ringkasan['sukses'] + ringkasan['gagal']
    ringkasan['tingkat_gagal_persen'] = (ringkasan['gagal'] / ringkasan['total'] * 100).round(2)
    return ringkasan.reset_index()


def simpan_ke_excel(ringkasan_performa, ringkasan_payload, ringkasan_asr, df_mentah):
    with pd.ExcelWriter(XLSX_PATH, engine='openpyxl') as writer:
        ringkasan_performa.to_excel(writer, sheet_name='Ringkasan Performa', index=False)
        ringkasan_payload.to_excel(writer, sheet_name='Ringkasan Payload', index=False)
        ringkasan_asr.to_excel(writer, sheet_name='Tingkat Kegagalan', index=False)
        df_mentah.to_excel(writer, sheet_name='Data Mentah', index=False)

    from openpyxl import load_workbook
    wb = load_workbook(XLSX_PATH)
    header_font = Font(name='Arial', bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')

    for nama_sheet in wb.sheetnames:
        ws = wb[nama_sheet]
        ws.freeze_panes = 'A2'
        for col_idx, cell in enumerate(ws[1], start=1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            panjang_maks = max(
                [len(str(cell.value))] + [len(str(r[col_idx - 1].value)) for r in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 50))]
            )
            ws.column_dimensions[get_column_letter(col_idx)].width = min(panjang_maks + 4, 40)
    wb.save(XLSX_PATH)
    print(f'✅ Tabel statistik tersimpan: {XLSX_PATH}')


def buat_grafik_durasi(ringkasan_performa):
    tahap_urutan = ringkasan_performa['tahap'].unique().tolist()
    fig, ax = plt.subplots(figsize=(10, 6))
    lebar = 0.35
    posisi = range(len(tahap_urutan))
    for i, metode in enumerate(['JWT', 'ZKP']):
        data_metode = ringkasan_performa[ringkasan_performa['metode'] == metode].set_index('tahap')
        rata_rata = [data_metode.loc[t, 'rata_rata'] if t in data_metode.index else 0 for t in tahap_urutan]
        std_dev = [data_metode.loc[t, 'std_dev'] if t in data_metode.index else 0 for t in tahap_urutan]
        offset = (i - 0.5) * lebar
        ax.bar([p + offset for p in posisi], rata_rata, lebar, yerr=std_dev, capsize=4,
               label=metode, color=WARNA.get(metode, 'gray'))
    ax.set_xticks(list(posisi))
    ax.set_xticklabels(tahap_urutan, rotation=20, ha='right')
    ax.set_ylabel('Durasi rata-rata (ms)')
    ax.set_title('Perbandingan Durasi Tiap Tahap: JWT vs ZKP Groth16')
    ax.legend(title='Metode')
    ax.grid(axis='y', linestyle='--', alpha=0.4)
    fig.tight_layout()
    fig.savefig(CHART_DURASI_PATH, dpi=150)
    plt.close(fig)
    print(f'✅ Grafik durasi tersimpan: {CHART_DURASI_PATH}')


def buat_boxplot_total_setup(df):
    data_setup = df[(df['tahap'] == 'call_setup_total') & (df['status'] == 'sukses') & df['durasi_ms'].notna()]
    if data_setup.empty:
        print('⚠️  Tidak ada data call_setup_total yang berhasil, boxplot dilewati.')
        return
    fig, ax = plt.subplots(figsize=(7, 6))
    kelompok = [data_setup[data_setup['metode'] == m]['durasi_ms'].values for m in ['JWT', 'ZKP']]
    kotak = ax.boxplot(kelompok, tick_labels=['JWT', 'ZKP'], patch_artist=True, showmeans=True)
    for patch, metode in zip(kotak['boxes'], ['JWT', 'ZKP']):
        patch.set_facecolor(WARNA.get(metode, 'gray'))
        patch.set_alpha(0.6)
    ax.set_ylabel('Total call setup (ms)')
    ax.set_title('Sebaran Waktu Total Call Setup: JWT vs ZKP Groth16')
    ax.grid(axis='y', linestyle='--', alpha=0.4)
    fig.tight_layout()
    fig.savefig(CHART_BOXPLOT_PATH, dpi=150)
    plt.close(fig)
    print(f'✅ Box plot tersimpan: {CHART_BOXPLOT_PATH}')


def buat_grafik_payload(ringkasan_payload):
    if ringkasan_payload.empty:
        print('⚠️  Tidak ada data payload, grafik payload dilewati.')
        return
    tahap_urutan = ringkasan_payload['tahap'].unique().tolist()
    fig, ax = plt.subplots(figsize=(9, 6))
    lebar = 0.35
    posisi = range(len(tahap_urutan))
    for i, metode in enumerate(['JWT', 'ZKP']):
        data_metode = ringkasan_payload[ringkasan_payload['metode'] == metode].set_index('tahap')
        rata_rata = [data_metode.loc[t, 'rata_rata'] if t in data_metode.index else 0 for t in tahap_urutan]
        offset = (i - 0.5) * lebar
        ax.bar([p + offset for p in posisi], rata_rata, lebar, label=metode, color=WARNA.get(metode, 'gray'))
    ax.set_xticks(list(posisi))
    ax.set_xticklabels(tahap_urutan, rotation=20, ha='right')
    ax.set_ylabel('Ukuran payload rata-rata (byte)')
    ax.set_title('Perbandingan Ukuran Payload: JWT vs ZKP Groth16')
    ax.legend(title='Metode')
    ax.grid(axis='y', linestyle='--', alpha=0.4)
    fig.tight_layout()
    fig.savefig(CHART_PAYLOAD_PATH, dpi=150)
    plt.close(fig)
    print(f'✅ Grafik payload tersimpan: {CHART_PAYLOAD_PATH}')


def main():
    print('📖 Membaca data/log-pengujian.csv ...')
    df = muat_data()
    print(f'   Total baris: {len(df)}')
    ringkasan_performa = hitung_ringkasan_performa(df)
    ringkasan_payload = hitung_ringkasan_payload(df)
    ringkasan_asr = hitung_asr(df)
    simpan_ke_excel(ringkasan_performa, ringkasan_payload, ringkasan_asr, df)
    buat_grafik_durasi(ringkasan_performa)
    buat_boxplot_total_setup(df)
    buat_grafik_payload(ringkasan_payload)
    print('\n📊 Ringkasan Performa (rata-rata durasi per tahap):')
    print(ringkasan_performa.to_string(index=False))
    print('\n📊 Tingkat Kegagalan per Tahap:')
    print(ringkasan_asr.to_string(index=False))
    print('\n✅ Semua tabel & grafik tersimpan di data/processed/ dan data/results/')


if __name__ == '__main__':
    main()